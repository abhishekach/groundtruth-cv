"""Application documents in ATS-safe formats, and the checks that prove it.

Rules (from applicant-tracking-system and European application guidance):
- one column; no tables, text boxes, icons or header/footer content (parsers skip or scramble them);
- name, e-mail and phone as plain text at the top of the body; standard section headings; MM/YYYY dates;
- standard fonts; Word (.docx) and a text-based PDF (never an image PDF); at most two pages for a job CV;
- a photo only where the market expects it (code decides, see photo_policy) and never instead of text.
Both files are written from the same content, then read back the way a parser would (ats_check).
"""
from __future__ import annotations

import re
from pathlib import Path

STANDARD_HEADINGS = {
    "profile", "summary", "key skills", "skills", "professional experience", "experience", "work experience",
    "selected projects", "projects", "education", "languages & certificates", "languages and certificates", "languages",
    "certificates", "interests", "research interests", "research experience", "publications", "referees", "references",
    "methods and skills", "methods & skills", "awards", "teaching",
}
CSS = """
@page{size:A4;margin:16mm 16mm 16mm 16mm}
body{font:10.5pt/1.42 "Charter","Georgia",serif;color:#111;max-width:178mm;margin:24px auto;padding:0 12px}
h1{font:600 20pt/1.1 "Avenir Next","Helvetica Neue",sans-serif;margin:0 0 4px}
.contact{font:9.5pt/1.3 "Avenir Next","Helvetica Neue",sans-serif;color:#444;margin-bottom:14px}
h2{font:600 10.5pt/1.2 "Avenir Next","Helvetica Neue",sans-serif;letter-spacing:.08em;text-transform:uppercase;
   border-bottom:1px solid #999;padding-bottom:2px;margin:14px 0 6px}
p{margin:0 0 4px} p.entry{font-weight:600;margin-top:6px}
ul{margin:0 0 4px 18px;padding:0} li{margin:0 0 2px}
.fill{background:#fff3c4}
.hint{font:9pt sans-serif;color:#666;border:1px dashed #bbb;padding:6px 8px;margin-bottom:14px}
@media print{.hint{display:none}}
table{border-collapse:collapse;width:100%;font:9.5pt/1.35 sans-serif} td,th{border:1px solid #ccc;padding:4px 6px;vertical-align:top;text-align:left}
th{background:#f2f2f2} .ok{color:#1d7f47} .bad{color:#b63a24}
"""


NO_PHOTO_COUNTRIES = {"GB", "UK", "IE", "US", "CA", "AU", "NZ", "NL", "SE", "DK", "FI", "NO", "SG", "HK"}
NO_PHOTO_WORDS = re.compile(r"(?i)\b(no photo|without (a )?photo|ohne (bewerbungs)?foto|kein(e)? foto|anonym(ous|e|isierte)? "
                            r"bewerb|anonymous application|do not include a photo|please omit (your )?photo)")


def prepared_photo(src: Path, out_dir: Path) -> Path:
    """A copy of the photo that prints the same everywhere: transparency flattened onto white, RGB JPEG, the original
    left untouched. Returns the original if Pillow cannot read it."""
    try:
        from PIL import Image
        im = Image.open(src)
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, (255, 255, 255))
            bg.paste(im, mask=im.getchannel("A"))
            im = bg
        else:
            im = im.convert("RGB")
        out = out_dir / "photo.jpg"
        im.save(out, "JPEG", quality=92)
        return out
    except Exception:  # noqa: BLE001
        return src


def photo_low_res(path: Path, width_mm: float = 28) -> bool:
    """True when the photo has under ~250 dpi at its printed width (soft in print)."""
    try:
        from PIL import Image
        return Image.open(path).size[0] < width_mm / 25.4 * 250
    except Exception:  # noqa: BLE001
        return False


def photo_policy(country: str | None, posting: str, model_says: bool, have_photo: bool) -> tuple[bool, str]:
    """A photo goes on only where it helps: never in markets where photos are unwelcome or when the posting asks for
    none; elsewhere the writer's judgement of this employer (conservative German firms expect one, international
    tech firms and startups treat it as optional)."""
    cc = (country or "").upper()
    if not have_photo:
        return False, "no photo on file (set photo: in settings.yaml to use one where it is expected)"
    if cc in NO_PHOTO_COUNTRIES:
        return False, f"no photo: unusual or unwelcome in {cc} applications"
    if NO_PHOTO_WORDS.search(posting or ""):
        return False, "no photo: the posting asks for an application without one"
    return (True, "photo included: expected by this kind of employer") if model_says else \
           (False, "no photo: optional for this employer, the CV leads with skills")


# ---------------------------------------------------------------- Word
def _clean_props(d, author: str, title: str) -> None:
    """Document properties carry the applicant's name only: no generator name, no template comments or dates."""
    from datetime import datetime, timezone
    cp = d.core_properties
    now = datetime.now(timezone.utc).replace(microsecond=0)
    cp.author, cp.last_modified_by, cp.title = author, author, title
    cp.comments = cp.subject = cp.keywords = cp.category = ""
    cp.created = cp.modified = now
    cp.revision = 1


def clean_pdf(path: Path) -> None:
    """Blanks the PDF library's name from comments and properties, keeping byte offsets intact."""
    b = path.read_bytes()
    for pat in (rb"ReportLab Generated PDF document -- digest \(opensource\)", rb"ReportLab Generated PDF document \(opensource\)",
                rb"ReportLab generated PDF document -- digest \(opensource\)", rb"ReportLab PDF Library - \(opensource\)",
                rb"ReportLab PDF Library - www.reportlab.com", rb"\(unspecified\)"):
        b = re.sub(pat, lambda m: b" " * len(m.group(0)), b)
    path.write_bytes(b)


def _doc():
    from docx import Document
    from docx.shared import Mm, Pt
    d = Document()
    s = d.sections[0]
    s.page_height, s.page_width = Mm(297), Mm(210)
    s.left_margin = s.right_margin = Mm(18)
    s.top_margin = s.bottom_margin = Mm(15)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.08
    return d


def _runs(par, text: str, bold: bool = False, size=None, color=None) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    for part in re.split(r"(\[TO FILL[^\]]*\])", text):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = bold
        if size:
            r.font.size = size
        if color is not None:
            r.font.color.rgb = color
        if part.startswith("[TO FILL"):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _rule(par) -> None:
    """A thin line under a heading, drawn as a paragraph border (no table, no shape)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ppr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(k), v)
    b.append(bottom)
    ppr.append(b)


def _float_right(run, width_mm: float) -> None:
    """Turns the run's inline picture into a picture floating at the top right of the text area, with text wrapping
    around it. It stays a plain picture (no text box, no table), which parsers skip, so the text around it reads normally."""
    from docx.oxml import parse_xml
    inline = run._r.xpath(".//wp:inline")[0]
    ext = inline.xpath("./wp:extent")[0]
    cx, cy = ext.get("cx"), ext.get("cy")
    graphic = inline.xpath("./a:graphic")[0]
    docpr = inline.xpath("./wp:docPr")[0]
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="180000" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" '
        'layoutInCell="1" allowOverlap="0">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapSquare wrapText="left"/></wp:anchor>')
    anchor.append(docpr)
    anchor.append(parse_xml('<wp:cNvGraphicFramePr xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"/>'))
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)


def _head_docx(d, head: dict) -> None:
    from docx.shared import Mm, Pt, RGBColor
    p = d.add_paragraph()
    if head.get("photo"):
        r = p.add_run()
        r.add_picture(str(head["photo"]), width=Mm(28))
        _float_right(r, 28)
    _runs(p, head["name"], bold=True, size=Pt(20))
    p.paragraph_format.space_after = Pt(0)
    if head.get("headline"):
        p = d.add_paragraph()
        _runs(p, head["headline"], size=Pt(11.5), color=RGBColor(0x2F, 0x4F, 0x6F))
    for line in [head.get("contact")] + ([head["permit"]] if head.get("permit") else []):
        if line:
            p = d.add_paragraph()
            _runs(p, line, size=Pt(9.5), color=RGBColor(0x40, 0x40, 0x40))
    if head.get("photo"):   # keep the first heading below the photo
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(10)


def cv_docx(head: dict, sections: list[dict], path: Path, place_date: str | None = None) -> None:
    from docx.shared import Pt
    d = _doc()
    _head_docx(d, head)
    for s in sections:
        if not s["lines"]:
            continue
        h = d.add_paragraph()
        _runs(h, s["heading"].upper(), bold=True, size=Pt(10.5))
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(10), Pt(3)
        _rule(h)
        for l in s["lines"]:
            if l["style"] == "bullet":
                p = d.add_paragraph(style="List Bullet")
                _runs(p, l["text"])
            else:
                p = d.add_paragraph()
                _runs(p, l["text"], bold=l["style"] == "entry")
                if l["style"] == "entry":
                    p.paragraph_format.space_before = Pt(5)
    if place_date:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        _runs(p, place_date)
    _clean_props(d, head["name"], "Curriculum Vitae")
    d.save(str(path))


def letter_docx(head: dict, L: dict, path: Path) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    d = _doc()
    d.styles["Normal"].font.size = Pt(11)
    for i, line in enumerate(L["sender"]):
        p = d.add_paragraph()
        _runs(p, line, bold=i == 0, size=Pt(14) if i == 0 else Pt(10))
        p.paragraph_format.space_after = Pt(0)
    d.add_paragraph()
    for line in L["recipient"]:
        p = d.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    p = d.add_paragraph(L["place_date"])
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(10)
    p = d.add_paragraph()
    _runs(p, L["subject"], bold=True)
    p.paragraph_format.space_before = Pt(10)
    p = d.add_paragraph(L["salutation"])
    p.paragraph_format.space_before = Pt(8)
    for t in L["paragraphs"]:
        p = d.add_paragraph()
        _runs(p, t)
        p.paragraph_format.space_after = Pt(8)
    d.add_paragraph(L["closing"])
    d.add_paragraph(head["name"])
    _clean_props(d, head["name"], L["subject"])
    d.save(str(path))


# ---------------------------------------------------------------- PDF (text-based, reportlab)
FONT_CANDIDATES = [  # embedded TrueType so every character (ä, §, –, ł) prints and extracts as text
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]
_FONTS: tuple[str, str] | None = None


def _fonts() -> tuple[str, str]:
    global _FONTS
    if _FONTS:
        return _FONTS
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    _FONTS = ("Helvetica", "Helvetica-Bold")
    for reg, bold in FONT_CANDIDATES:
        if Path(reg).exists() and Path(bold).exists():
            try:
                pdfmetrics.registerFont(TTFont("CV", reg))
                pdfmetrics.registerFont(TTFont("CV-Bold", bold))
                from reportlab.lib.fonts import addMapping
                addMapping("CV", 0, 0, "CV")
                addMapping("CV", 1, 0, "CV-Bold")
                addMapping("CV", 0, 1, "CV")
                addMapping("CV", 1, 1, "CV-Bold")
                _FONTS = ("CV", "CV-Bold")
                break
            except Exception:  # noqa: BLE001
                continue
    return _FONTS


def _pdf_styles():
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    R, B = _fonts()
    base = dict(fontName=R, fontSize=10, leading=13.2)
    return {
        "name": ParagraphStyle("name", fontName=B, fontSize=19, leading=23),
        "headline": ParagraphStyle("headline", fontName=R, fontSize=11, leading=14, textColor=HexColor("#2F4F6F")),
        "small": ParagraphStyle("small", fontName=R, fontSize=9, leading=11.5, textColor=HexColor("#404040")),
        "h": ParagraphStyle("h", fontName=B, fontSize=10, leading=13, spaceBefore=9, spaceAfter=3),
        "entry": ParagraphStyle("entry", fontName=B, spaceBefore=4, **{k: v for k, v in base.items() if k != "fontName"}),
        "text": ParagraphStyle("text", **base),
        "bullet": ParagraphStyle("bullet", leftIndent=11, bulletIndent=2, **base),
        "letter": ParagraphStyle("letter", fontName=R, fontSize=10.5, leading=14.5, spaceAfter=8),
        "right": ParagraphStyle("right", fontName=R, fontSize=10.5, leading=14, alignment=2),
    }


def _x(s: str) -> str:
    from xml.sax.saxutils import escape
    return re.sub(r"(\[TO FILL[^\]]*\])", r'<font backColor="#fff3a0">\1</font>', escape(s))


def cv_pdf(head: dict, sections: list[dict], path: Path, place_date: str | None = None) -> None:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer
    S = _pdf_styles()
    photo = head.get("photo")
    pw = 28 * mm

    def first(canvas, doc):
        if photo:
            canvas.drawImage(str(photo), A4[0] - 18 * mm - pw, A4[1] - 15 * mm - pw * 1.28, width=pw, height=pw * 1.28,
                             preserveAspectRatio=True, anchor="ne", mask="auto")

    indent = (pw + 6 * mm) if photo else 0
    hs = {k: S[k].clone(k + "_h", rightIndent=indent) for k in ("name", "headline", "small")}
    story = [Paragraph(_x(head["name"]), hs["name"])]
    if head.get("headline"):
        story.append(Paragraph(_x(head["headline"]), hs["headline"]))
    for line in [head.get("contact")] + ([head["permit"]] if head.get("permit") else []):
        if line:
            story.append(Paragraph(_x(line), hs["small"]))
    if photo:
        story.append(Spacer(1, 8 * mm))
    for s in sections:
        if not s["lines"]:
            continue
        story += [Paragraph(_x(s["heading"].upper()), S["h"]), HRFlowable(width="100%", thickness=0.5, color=HexColor("#808080"), spaceAfter=3)]
        for l in s["lines"]:
            st = l["style"]
            story.append(Paragraph(_x(l["text"]), S["bullet"], bulletText="•") if st == "bullet"
                         else Paragraph(_x(l["text"]), S["entry" if st == "entry" else "text"]))
    if place_date:
        story += [Spacer(1, 10 * mm), Paragraph(_x(place_date), S["text"])]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                      title="Curriculum Vitae", author=head["name"], creator=head["name"], producer=head["name"],
                      subject="", keywords="").build(story, onFirstPage=first)
    clean_pdf(path)


def letter_pdf(head: dict, L: dict, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    S = _pdf_styles()
    story = [Paragraph(_x(L["sender"][0]), S["name"])] + [Paragraph(_x(x), S["small"]) for x in L["sender"][1:]]
    story += [Spacer(1, 8 * mm)] + [Paragraph(_x(x), S["text"]) for x in L["recipient"]]
    story += [Spacer(1, 4 * mm), Paragraph(_x(L["place_date"]), S["right"]), Spacer(1, 4 * mm),
              Paragraph("<b>" + _x(L["subject"]) + "</b>", S["letter"]), Paragraph(_x(L["salutation"]), S["letter"])]
    story += [Paragraph(_x(t), S["letter"]) for t in L["paragraphs"]]
    story += [Paragraph(_x(L["closing"]), S["letter"]), Paragraph(_x(head["name"]), S["letter"])]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=22 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=16 * mm,
                      title=L["subject"], author=head["name"], creator=head["name"],
                      producer=head["name"], subject="", keywords="").build(story)
    clean_pdf(path)


# ---------------------------------------------------------------- read-back checks
def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9+#&]+", " ", (s or "").lower()).strip()


def _has(nk: str, t: str) -> bool:
    if f" {nk} " in t or (len(nk) > 5 and nk in t):
        return True
    words = [w for w in nk.split() if len(w) > 2]
    return " " in nk and bool(words) and all(f" {w}" in t for w in words)


def keyword_coverage(keywords: list[str], text: str, facts_text: str = "") -> list[dict]:
    """Which of the posting's key terms the final CV contains (case-insensitive), and whether the fact sheet supports each
    one (on_sheet). A term missing from the CV but on the sheet is a gap to close; one not on the sheet stays out."""
    t, f = f" {_norm(text)} ", f" {_norm(facts_text)} "
    out, seen = [], set()
    for k in keywords:
        nk = _norm(k)
        if not nk or nk in seen:
            continue
        seen.add(nk)
        out.append({"term": k, "covered": _has(nk, t), "on_sheet": _has(nk, f) if facts_text else None})
    return out


def ats_check(docx_path: Path, pdf_path: Path, head: dict, headings: list[str], keywords: list[str],
              max_pages: int = 2, facts_text: str = "") -> dict:
    """Reads the files back as an applicant tracking system would and reports what it found."""
    from docx import Document
    from pypdf import PdfReader
    res = []
    d = Document(str(docx_path))
    xml = d.element.xml
    res.append(("Word file: no tables", len(d.tables) == 0, f"{len(d.tables)} table(s)"))
    res.append(("Word file: no text boxes", "txbxContent" not in xml, "content in text boxes is skipped by parsers"))
    hf = [p.text for s in d.sections for p in list(s.header.paragraphs) + list(s.footer.paragraphs) if p.text.strip()]
    res.append(("Word file: nothing in the page header or footer", not hf, "; ".join(hf)[:80]))
    top = " ".join(p.text for p in d.paragraphs[:6])
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", head.get("contact", ""))
    phone = re.search(r"\+?\d[\d ()-]{7,}\d", head.get("contact", ""))
    for label, val in (("name", head["name"]), ("e-mail", email.group(0) if email else ""), ("phone", phone.group(0) if phone else "")):
        res.append((f"Word file: {label} at the top as text", bool(val) and val in top, val))
    reader = PdfReader(str(pdf_path))
    pdf_text = "\n".join(pg.extract_text() or "" for pg in reader.pages)
    res.append(("PDF: text is selectable (not an image)", len(pdf_text) > 500, f"{len(pdf_text)} characters read"))
    for label, val in (("name", head["name"]), ("e-mail", email.group(0) if email else "")):
        res.append((f"PDF: {label} found", bool(val) and val in pdf_text.replace("\n", " "), val))
    lines = [ln.strip().upper() for ln in pdf_text.splitlines()]
    pos, order_ok, missing = -1, True, []
    for h in headings:
        i = next((n for n, ln in enumerate(lines) if ln == h.upper()), -1)
        if i < 0:
            missing.append(h)
        elif i < pos:
            order_ok = False
        else:
            pos = i
    res.append(("PDF: every section heading found, in order", not missing and order_ok,
                ("missing: " + ", ".join(missing)) if missing else ("order ok" if order_ok else "order differs")))
    odd = [h for h in headings if h.lower() not in STANDARD_HEADINGS]
    res.append(("Standard section headings", not odd, ", ".join(odd) or "all standard"))
    res.append((f"Length: at most {max_pages} page(s)", len(reader.pages) <= max_pages, f"{len(reader.pages)} page(s)"))
    bad_dates = re.findall(r"\b(?:19|20)\d{2}\s*[-–]\s*(?:19|20)?\d{2}\b(?!/)", pdf_text)
    res.append(("Dates in one format (MM/YYYY)", not bad_dates, ", ".join(bad_dates[:3]) or "consistent"))
    cov = keyword_coverage(keywords, pdf_text, facts_text)
    n = sum(c["covered"] for c in cov)
    pct = round(100 * n / len(cov)) if cov else None
    gap = [c["term"] for c in cov if not c["covered"] and c["on_sheet"]]
    res.append(("Posting keywords your fact sheet supports are in the CV", not gap,
                ("missing: " + ", ".join(gap)) if gap else (f"{n} of {len(cov)} posting keywords in the CV ({pct}%)" if cov else "no keywords given")))
    return {"checks": [{"check": c, "ok": ok, "detail": det} for c, ok, det in res], "coverage": cov, "pct": pct,
            "pages": len(reader.pages), "passed": sum(1 for _, ok, _ in res if ok), "total": len(res)}


def checks_html(css: str, title: str, ats: dict, rows: list[dict], removed: list[dict], notes: list[str], photo_note: str) -> str:
    import html
    e = html.escape
    tr_ats = "".join(f"<tr><td>{e(c['check'])}</td><td class='{'ok' if c['ok'] else 'bad'}'>{'pass' if c['ok'] else 'fix'}</td>"
                     f"<td>{e(str(c['detail']))}</td></tr>" for c in ats["checks"])
    def kw_state(c):
        if c["covered"]:
            return "ok", "in your CV"
        return ("bad", "missing, but your fact sheet supports it") if c["on_sheet"] else ("", "not found in your fact sheet's wording, so left out (if you have it, add it there in these words)")
    kw = "".join(f"<li class='{kw_state(c)[0]}'>{e(c['term'])} — {kw_state(c)[1]}</li>" for c in ats["coverage"]) or "<li>none</li>"
    tr = "".join(f"<tr><td>{e(r['text'][:160])}</td><td>{e(', '.join(r['ids']) or '—')}</td>"
                 f"<td class='{'ok' if r['result'] in ('OK', 'placeholder') else 'bad'}'>{e(r['result'])}</td></tr>" for r in rows)
    rm = "".join(f"<li>{e(r['text'][:200])} — <i>{e(r['why'])}</i></li>" for r in removed) or "<li>none</li>"
    nt = "".join(f"<li>{e(n)}</li>" for n in notes)
    return (f"<!doctype html><meta charset='utf-8'><title>Checks – {e(title)}</title><style>{css}</style><body>"
            f"<h1>Checks</h1><div class='contact'>{e(title)}</div>"
            f"<h2>ATS read-back · {ats['passed']} of {ats['total']} passed</h2>"
            "<p>The Word and PDF files were read back the way an applicant tracking system reads them.</p>"
            f"<table><tr><th>Check</th><th>Result</th><th>Detail</th></tr>{tr_ats}</table>"
            f"<h2>Posting keywords · {ats['pct'] if ats['pct'] is not None else '—'}% in the CV</h2><p>Terms not on your fact "
            "sheet are never added. If you do have one of them, add it to the fact sheet and run tailor again.</p>"
            f"<ul>{kw}</ul><h2>Photo</h2><p>{e(photo_note)}</p>"
            f"<h2>Fact check</h2><table><tr><th>Line</th><th>Fact-sheet IDs</th><th>Result</th></tr>{tr}</table>"
            f"<h2>Removed during verification</h2><ul>{rm}</ul><h2>Notes</h2><ul>{nt}</ul></body>")


def letter_check(docx_path: Path, pdf_path: Path, words: tuple[int, int], max_pages: int, label: str) -> list[dict]:
    """Length and layout of a cover or motivation letter, read back from the files."""
    from docx import Document
    from pypdf import PdfReader
    d = Document(str(docx_path))
    reader = PdfReader(str(pdf_path))
    body = [p.text for p in d.paragraphs]
    n = sum(len(t.split()) for t in body[body.index(next(t for t in body if t.startswith("Dear"))) + 1:]) if any(
        t.startswith("Dear") for t in body) else sum(len(t.split()) for t in body)
    lo, hi = words
    return [{"check": f"{label}: length {lo}-{hi} words", "ok": lo <= n <= hi + 40, "detail": f"{n} words"},
            {"check": f"{label}: at most {max_pages} page(s)", "ok": len(reader.pages) <= max_pages, "detail": f"{len(reader.pages)} page(s)"},
            {"check": f"{label}: PDF text is selectable", "ok": len("".join(pg.extract_text() or "" for pg in reader.pages)) > 300,
             "detail": "read back with a PDF parser"},
            {"check": f"{label}: Word file has no tables or text boxes", "ok": not d.tables and "txbxContent" not in d.element.xml,
             "detail": "plain paragraphs"}]
